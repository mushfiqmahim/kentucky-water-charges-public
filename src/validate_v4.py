"""Numerical checks: arithmetic, membership, documentary counts and document tables.

These checks reuse source coding; they do not constitute independent source reading.
"""
from pathlib import Path
from decimal import Decimal
import hashlib, json, csv, math
import numpy as np
import pandas as pd
from docx import Document
from check_inputs import validate_inputs, load_reference

R = Path(__file__).resolve().parents[1]
F = R / 'data/reference'
def read(path):
    return pd.read_csv(path, dtype={'utility_id': str})
def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()
def cents(x):
    return int(Decimal(str(x)).quantize(Decimal('.01')) * 100)

# Fixed current expectations were captured before execution; no build regenerates them.
catalog_by_path = validate_inputs(R)
reference = load_reference(R)
def equal_numeric(a,b,where):
    if isinstance(a,bool) or isinstance(b,bool):
        assert type(a) is type(b) and a == b, where
    elif isinstance(a,int):
        assert isinstance(b,int) and a==b, (where,a,b)
    elif isinstance(a,(int,float)) and isinstance(b,(int,float)):
        assert math.isfinite(a) and math.isfinite(b), where
        assert abs(a-b) <= reference['numeric_absolute_tolerance'], (where,a,b)
    elif isinstance(a,dict):
        assert isinstance(b,dict) and a.keys()==b.keys(), where
        for k in a: equal_numeric(a[k],b[k],where+'.'+k)
    elif isinstance(a,list):
        assert isinstance(b,list) and len(a)==len(b), where
        for i,(x,y) in enumerate(zip(a,b)): equal_numeric(x,y,where+f'[{i}]')
    else: assert a == b, (where,a,b)
for path,digest in reference['inputs_sha256'].items():
    assert sha(R/path)==digest, ('authored input changed',path)
for name,expected in reference['json'].items():
    equal_numeric(expected,json.loads((R/'results'/name).read_text()),name)
for name,expected in reference['csv'].items():
    actual=list(csv.reader((R/'results'/name).open()))
    assert len(expected)==len(actual) and expected[0]==actual[0], name
    for rownum,(a,b) in enumerate(zip(expected[1:],actual[1:]),1):
        assert len(a)==len(b), (name,rownum)
        for col,x,y in zip(expected[0],a,b):
            if x==y: continue
            exact_columns={'utility_id','district_id','event_id','year','2019','2020','2021','2022','2023','2024','gallons','n','df','decline_n','neutral_n','growth_n','arc_all','arc_mixed','merger_flag','conservative_screen_pass','main_eligible','account_screen_pass','primary','endpoint_pair_reconstructed'}
            assert col not in exact_columns, (name,rownum,col,'exact value mismatch',x,y)
            try: xx,yy=float(x),float(y)
            except ValueError: raise AssertionError((name,rownum,col,x,y))
            assert math.isfinite(xx) and math.isfinite(yy), (name,rownum,col)
            assert abs(xx-yy)<=reference['numeric_absolute_tolerance'], (name,rownum,col,x,y)
            if col.startswith(('base_bill_','charge_inclusive_bill_','base_2019','base_2024','total_2019','total_2024','extra_2019','extra_2024','delta_base','delta_extra','delta_total')):
                assert cents(x)==cents(y), (name,rownum,col,'cent mismatch')
v4=json.loads((R/'results/models_v4.json').read_text())
assert len(v4)==13 and sum(len(x['coefficients']) for x in v4)==41
p=read(R/'results/analysis_panel_v4.csv')
s=p[p.gallons.eq(4000) & p.main_eligible].set_index('utility_id')
assert len(s)==20 and len(p[p.gallons.eq(4000) & p.price_status.eq('supported')])==27
for name,digest in reference['figures_sha256'].items():
    assert sha(R/'results/figures'/name)==digest, ('figure bytes',name)
# Declared source identities remain fixed; opening originals is a separate command.
catalog=read(R/'sources/catalog.csv')
assert len(catalog)==213 and catalog.path.nunique()==213
for row in catalog.itertuples():
    fixed=reference['sources'][row.path]
    assert row.sha256==fixed['sha256'] and row.git_blob==fixed['git_blob'] and int(row.bytes)==int(fixed['bytes']), ('fixed source identity',row.path)
source_ids={sid for cell in catalog.source_ids for sid in json.loads(cell)}
def check_source_references(value):
    if isinstance(value,dict):
        for key,v in value.items():
            if key=='source_id': assert v in source_ids, ('source_id',v)
            check_source_references(v)
    elif isinstance(value,list):
        for v in value: check_source_references(v)
    elif isinstance(value,str) and value.startswith('sources/'):
        assert value in catalog_by_path, ('source path',value)
for name in ['events_current.json','surcharges.json','coding_provenance.json']:
    check_source_references(json.loads((R/'data'/name).read_text()))
for path in (R/'data').glob('*.csv'):
    for row in csv.DictReader(path.open()):
        for value in row.values(): check_source_references(value)

# Endpoint and account coding match fixed declared source identities.
reg = read(R/'data/Source_Verification_Register_v4.csv')
assert len(reg) == 40 and len(reg[['utility_id','year']].drop_duplicates()) == 40
for x in reg.itertuples():
    assert catalog_by_path[x.source_file]['sha256'] == x.source_sha256
accounts = read(R/'data/Endpoint_Account_Source_Check.csv')
assert len(accounts) == 40
for x in accounts.itertuples():
    assert catalog_by_path[x.source_file]['sha256'] == x.sha256
    assert s.loc[x.utility_id, str(x.year)] == x.residential_end

# Every arithmetic row is checked against current analytical outputs at cent precision.
rec = read(R/'results/Reconstruction_Summary_20_v4.csv')
assert len(rec) == rec.utility_id.nunique() == 20 and set(rec.utility_id) == set(s.index)
for x in rec.itertuples():
    a = s.loc[x.utility_id]
    for yr in [2019, 2024]:
        assert cents(getattr(x,f'base_{yr}')) == cents(a[f'base_bill_{yr}'])
        assert cents(getattr(x,f'total_{yr}')) == cents(a[f'charge_inclusive_bill_or_scenario_{yr}'])
        assert cents(getattr(x,f'base_{yr}')) + cents(getattr(x,f'extra_{yr}')) == cents(getattr(x,f'total_{yr}'))
    for kind in ['base', 'extra', 'total']:
        assert cents(getattr(x,f'{kind}_2024')) - cents(getattr(x,f'{kind}_2019')) == cents(getattr(x,f'delta_{kind}'))
assert [sum(map(cents,rec[c])) for c in ['delta_base','delta_extra','delta_total']] == [20714,733,21447]
assert (int((rec.delta_base>0).sum()),int((rec.delta_total==0).sum()),int((rec.delta_extra>0).sum())) == (18,2,4)
assert 'Dec 2020' in rec.set_index('utility_id').loc['23300','temporary_extra_absent_both_endpoints']

# Dispositions preserve both price support and the original account screen.
disp = read(R/'results/Selected_34_Disposition_v4.csv').set_index('utility_id')
orig = read(R/'data/selection_frame_34.csv').set_index('utility_id')
q = p[p.gallons.eq(4000)].set_index('utility_id')
assert len(disp) == len(disp.index.unique()) == 34 and set(disp.index) == set(orig.index)
assert len(q) == 32 and set(disp[disp.endpoint_pair_reconstructed].index) == set(q.index)
assert set(disp[disp.primary].index) == set(s.index)
assert disp.price_classification.eq('Supported').sum() == 27
assert disp.price_classification.eq('Conditional').sum() == 5
assert disp.original_stratum.value_counts().to_dict() == {'nondeclining':17,'declining':17}
assert disp[disp.primary].original_stratum.value_counts().to_dict() == {'nondeclining':11,'declining':9}
for uid,x in disp.iterrows():
    assert x.original_stratum == orig.loc[uid,'original_stratum']
    assert x.account_screen_pass == orig.loc[uid,'conservative_screen_pass']
    assert bool(x.primary) == bool(x.price_classification == 'Supported' and x.account_screen_pass)
assert set(disp[disp.price_classification.eq('Supported') & ~disp.account_screen_pass].index) == {'21200','31200','26600','31800','23000','18500','19201'}

# Current coded events are inputs. E057 retains its prior label and correction provenance.
old=json.loads((R/'data/events_current.json').read_text())
new=json.loads((R/'results/events_v4.json').read_text())
assert len(new)==len(old)==76 and old==new
source_hashes=set()
for event in new:
    if event['event_id']=='E057':
        assert 'wholesale special-contract' in event['regulatory_category']
        assert event['regulatory_category_v3'] and event['v4_classification_note']
    for src in event['resolved_sources']:
        assert catalog_by_path[src['path']]['sha256']==src['sha256']
        source_hashes.add(src['sha256'])
l = pd.read_csv(R/'results/Event_Ledger_Classification_v4.csv')
schedules = l[~l.record_type.eq('amending_order')]
surcharges = json.loads((R/'data/surcharges.json').read_text())
cases = {x['case'] for x in new+surcharges if x.get('case')}
actual = dict(schedule_authorization_rows=len(schedules), baseline_anchors=int(l.record_type.eq('baseline_anchor').sum()),
    subsequent_schedule_authorization_rows=int((~schedules.record_type.eq('baseline_anchor')).sum()),
    arithmetic_price_change_rows=int(l.base_price_changed_4000.eq(True).sum()),
    zero_base_change_rows=int(l.base_price_changed_4000.eq(False).sum()),
    amendment_rows=int(l.record_type.eq('amending_order').sum()),
    superseded_schedule_rows=int(l.record_type.eq('superseded_schedule').sum()),
    separate_recurring_charge_records=len(surcharges),combined_logical_records=len(new)+len(surcharges),
    distinct_recorded_case_ids=len(cases), schedule_rows_without_case_id=int(schedules['case'].isna().sum()))
counts = json.loads((R/'results/ledger_counts_v4.json').read_text())
for k,v in actual.items(): assert counts[k] == v, (k,v,counts[k])
assert counts['legally_operative_period_count'] is None
assert counts['unique_implemented_price_change_count'] is None

# Knott's internal arithmetic is distinct from the unreconciled summary assertion.
assert 2908014 + 47791 + 9558 == 2965363
assert 2965363 - (33930+975+8846+1250850) == 1670762
assert 1670762 - 992944 == 677818
assert 677818 - 646151 == 31667
for uid,expected in [('19400',[3.7643208,68.7563538]),('23300',[-1.1723329,86.8170267])]:
    a=s.loc[uid]
    assert np.allclose([100*(a['2024']/a['2019']-1),a.full_pct_change],expected,rtol=0,atol=1e-7)

# Compare every rendered DOCX table to generated, data-linked text.
t=json.loads((R/'results/manuscript_tables.json').read_text())
for key,expected in reference['manuscript_tables'].items():
    assert t[key][1:3]==expected[1:3], ('fixed manuscript table',key)
doc=Document(R/'paper/KEA_Manuscript_v4.docx');assert len(doc.tables)==7
for table,key in zip(doc.tables,['TABLE1','TABLE2','TABLE3','TABLE4','TABLEA1','TABLEB1','TABLEB2']):
    title,headers,rows,note=t[key]
    assert [[c.text for c in row.cells] for row in table.rows] == [headers]+[[str(c) for c in row] for row in rows]
st=json.loads((R/'results/supplement_tables_v4.json').read_text())
for key,expected in reference['supplement_tables'].items():
    assert st[key][1:3]==expected[1:3], ('fixed supplement table',key)
sd=Document(R/'paper/Supplement_v4.docx');assert len(sd.tables)==5
for table,key in zip(sd.tables,['S1','S2','S3','S4']):
    title,headers,rows,widths=st[key]
    assert [[c.text for c in row.cells] for row in table.rows] == [headers]+[[str(c) for c in row] for row in rows]
financial=[line for line in (R/'paper/Source_Followup_v4.md').read_text().splitlines() if line.startswith('|')]
expected=[[x.strip() for x in line.strip('|').split('|')] for i,line in enumerate(financial) if i!=1]
assert expected==reference['financial_table'], 'fixed financial table'
assert [[c.text for c in row.cells] for row in sd.tables[4].rows] == expected
md=(R/'paper/manuscript_complete.md').read_text()
assert '{{' not in md
for token in ['−1.231','[−2.794, 0.332]','68 schedule','43 distinct','24.24%','207.14','7.33','214.47']:
    assert token in md,token

result={'status':'passed','primary_n':20,'supported_n':27,'selected_n':34,'reconstructed_n':32,
 'fixed_reference_models_checked':13,'coefficient_records_checked':41,'unchanged_analytical_inputs':True,'unchanged_figure_pngs':2,'declared_source_identities_checked':213,'original_source_files_checked':False,'numeric_absolute_tolerance':1e-9,'numeric_relative_tolerance':0,'descriptive_absolute_tolerance':1e-7,'categorical_and_cents':'exact',
 'endpoint_declared_source_rows_checked':40,'endpoint_account_coding_rows_checked':40,
 'distinct_declared_event_source_hashes_checked':len(source_hashes),'reconstruction_rows_checked':20,
 'disposition_rows_checked':34,'manuscript_tables_checked':7,'supplement_tables_checked':5,
 'component_totals_dollars':[207.14,7.33,214.47],'ledger_counts_recomputed':actual,
 'knott_summary_discrepancy_dollars_unresolved':31667,
 'scope':'Numerical reproduction from compiled inputs and coded evidence; original source-file verification is separate. Neither workflow independently validates documentary interpretation or reconstructs original collection and screening.'}
(R/'results/numerical_validation_v4.json').write_text(json.dumps(result,indent=2))
print(json.dumps(result,indent=2))
