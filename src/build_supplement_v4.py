from pathlib import Path
import pandas as pd,json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
disp=pd.read_csv(R/'results/Selected_34_Disposition_v4.csv',dtype={'utility_id':str});rec=pd.read_csv(R/'results/Reconstruction_Summary_20_v4.csv',dtype={'utility_id':str})
short=lambda n:n.replace(' County Water and Sewer District',' Co.').replace(' County Water & Sewer District',' Co.').replace(' County Water District',' Co.').replace(' Water District','').replace('Cumberland Falls Highway','Cumberland Falls Hwy.').replace('Garrison-Quincy-Ky-O-Heights','Garrison').replace('Southern Water & Sewer District','Southern Water & Sewer')
T={}
T['S1']=('Table S1. Complete disposition of the 34 selected districts',['District and ID','Original stratum','Price pair','Account screen','Primary','Decision and evidence'],[[short(v.utility_name)+' ('+v.utility_id+')','Decline' if v.original_stratum=='declining' else 'Nondecline',v.price_classification if v.endpoint_pair_reconstructed else 'No pair','Pass' if v.account_screen_pass else 'Fail','Yes' if v.primary else 'No',v.decision_reason+' Source: '+v.evidence_reference+'.'] for v in disp.itertuples()],[1.25,.85,.85,.65,.65,2.25])
n=pd.read_csv(R/'results/neutral_band_diagnostics.csv');n=n[n.version.eq('v4')]
T['S2']=('Table S2. Current neutral-band descriptions',['Band %','Decline n','Neutral n','Growth n','Decline median %','Growth median %','Gap pp'],[[f'{v.band_percent:g}',str(v.decline_n),str(v.neutral_n),str(v.growth_n),f'{v.decline_median:.4f}',f'{v.growth_median:.4f}',f'{v.gap_pp:.4f}'] for v in n.itertuples()],[.60,.70,.70,.70,1.25,1.25,1.30])
g=pd.read_csv(R/'results/geographic_omission_diagnostics.csv');g=g[g.version.eq('v4')]
T['S3']=('Table S3. Current geographic diagnostics',['Sample','N','Slope','95% HC3 interval'],[[v.model,str(v.n),f'{v.slope:.6f}',f'[{v.ci_low:.6f}, {v.ci_high:.6f}]'] for v in g.itertuples()],[2.3,.5,1.2,2.5])
b=pd.read_csv(R/'results/models_v4.csv')
T['S4']=('Table S4. Retained current specifications',['Specification','N','Slope','95% HC3 interval'],[[v.model,str(v.n),f'{v.slope:.6f}',f'[{v.ci_low:.6f}, {v.ci_high:.6f}]'] for v in b.itertuples()],[2.5,.35,1.25,2.40])
fragment=re.sub(r'^#{1,2} (.+)$',r'### \1',(R/'paper/Source_Followup_v4.md').read_text(),flags=re.MULTILINE)
body=(R/'paper/supplement.md').read_text().replace('{{SOURCE_FOLLOWUP}}',fragment)
# Financial markdown table is rendered natively along with placeholders.
def mdtable(headers,rows):return '| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(str(v).replace('|','/') for v in row)+' |' for row in rows)
complete=body
for key,(title,h,rows,w) in T.items():complete=complete.replace('{{'+key+'}}',title+'\n\n'+mdtable(h,rows))
(R/'paper/Supplement_v4.md').write_text(complete)
# Standalone readable versions reuse the same generated records, not independently typed tables.
t=json.loads((R/'results/manuscript_tables.json').read_text())
(R/'results/Reconstruction_Summary_20_v4.md').write_text('# Full sample reconstruction\n\n'+ '\n\n'.join(v[0]+'\n\n'+mdtable(v[1],v[2])+'\n\n'+v[3] for k,v in t.items() if k in ['TABLEB1','TABLEB2']))
(R/'results/Selected_34_Disposition_v4.md').write_text('# Complete selected sample disposition\n\n'+mdtable(T['S1'][1],T['S1'][2])+'\n\n34 selected; 32 reconstructed; 27 supported; 20 primary. Pass refers to the stated mechanical account screen, not proven physical-territory stability.\n')
d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
st=d.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(11);st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(7)
for nm,sz in [('Title',17),('Heading 1',13),('Heading 2',11)]:d.styles[nm].font.name='Times New Roman';d.styles[nm].font.size=Pt(sz);d.styles[nm].font.color.rgb=RGBColor(0,0,0)
for sty in d.styles:
 if sty.type==1:
  pr=sty.element.find(qn('w:pPr'))
  if pr is not None:
   for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
  rp=sty.element.find(qn('w:rPr'))
  if rp is not None:
   rf=rp.find(qn('w:rFonts'))
   if rf is not None:
    for attr in list(rf.attrib):
     if 'Theme' in attr:del rf.attrib[attr]
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
def addtable(title,headers,rows,widths):
 if title:
  p=d.add_paragraph(title);p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for ci,col in enumerate(t.columns):col.width=Inches(widths[ci])
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    no=OxmlElement('w:'+side);no.set(qn('w:w'),'30' if title.startswith('Table S4') else '60');no.set(qn('w:type'),'dxa');mar.append(no)
   pr.append(mar)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');pr.append(sh)
   for pp in c.paragraphs:
    pp.paragraph_format.line_spacing=1;pp.paragraph_format.space_before=Pt(3 if title.startswith('Table S4') else 4);pp.paragraph_format.space_after=Pt(3 if title.startswith('Table S4') else 4);pp.paragraph_format.keep_with_next=(ri==0 or not title)
    for rr in pp.runs:rr.font.size=Pt(10);rr.bold=ri==0
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  no=OxmlElement('w:'+side);no.set(qn('w:val'),'single');no.set(qn('w:sz'),'4');no.set(qn('w:color'),'D9D9D9');borders.append(no)
 t._tbl.tblPr.append(borders)
 d.add_paragraph().paragraph_format.space_after=Pt(2)
for block in body.split('\n\n'):
 block=block.strip()
 if not block:continue
 key=block.removeprefix('{{').removesuffix('}}')
 if key in T:addtable(*T[key])
 elif block.startswith('|'):
  lines=block.splitlines();h=[x.strip() for x in lines[0].strip('|').split('|')];rows=[[x.strip() for x in l.strip('|').split('|')] for l in lines[2:]];addtable('',h,rows,[2.3,1.0,3.2])
 elif block.startswith('# '):d.add_paragraph(block[2:],'Title')
 elif block.startswith('## '):d.add_paragraph(block[3:],'Heading 1')
 elif block.startswith('### '):d.add_paragraph(block[4:],'Heading 2')
 else:
  block=re.sub(r'\[([^]]+)\]\((https?://[^)]+)\)',r'\1 (\2)',block).replace('`','')
  d.add_paragraph(block)
d.save(R/'paper/Supplement_v4.docx');(R/'results/supplement_tables_v4.json').write_text(json.dumps(T,indent=2));print('Supplement and standalone readable tables built')
