from pathlib import Path
import json,re
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
R=Path(__file__).resolve().parents[1];E=R/'data';d=Document();s=d.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.space_after=Pt(6)
for k,size in [('Title',19),('Heading 1',14),('Heading 2',12)]:
 st=d.styles[k];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6)
for st in d.styles:
 if st.type==1:
  pr=st.element.find(qn('w:pPr'))
  if pr is not None:
   for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
  rp=st.element.find(qn('w:rPr'))
  if rp is not None:
   rf=rp.find(qn('w:rFonts'))
   if rf is not None:
    for a in list(rf.attrib):
     if 'Theme' in a:del rf.attrib[a]
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
d.core_properties.author='MD. Mushfiquzzaman Mahim';d.core_properties.title='Customer Base Change and Approved Water Charges in Kentucky 2019 to 2024'
md=(R/'paper/manuscript.md').read_text()
tables={}
m=pd.read_csv(R/'results/primary_districts_v4.csv',dtype={'utility_id':str});fits={x['model']:x for x in json.loads((R/'results/models_v4.json').read_text())}
tables['TABLE1']=('Table 1. Sample construction',['Stage','Districts','Initial decliners'],[['Original matched-name candidates','97','17'],['Historical tariff draw: all initial decliners + 17 sampled nondecliners','34','17'],['Reconstructed price pairs, including conditional cases','32','16'],['Supported benchmark price pairs','27','14'],['Primary: supported prices and account screen','20','9']], 'Daviess merger exclusion is a separate frame correction (97 to 96); Daviess was not in the 34-district draw. Initial classifications preceded tariff review. The primary sample is selected, not a representative statewide panel.')
m['account_pct']=100*(m['2024']/m['2019']-1)
cols=[('Residential accounts, 2019','2019'),('Account change, ordinary percent','account_pct'),('2019 monthly charge, nominal dollars','charge_inclusive_bill_or_scenario_2019'),('2024 monthly charge, nominal dollars','charge_inclusive_bill_or_scenario_2024'),('Nominal charge growth, percent','full_pct_change'),('Real charge growth, percent','real_pct_change')]
rows=[]
for label,c in cols:
 v=m[c];rows.append([label]+[f'{x:,.2f}' for x in [v.mean(),v.std(),v.min(),v.median(),v.max()]])
tables['TABLE2']=('Table 2. Main-sample descriptive statistics',['Variable','Mean','SD','Min','Median','Max'],rows,'N = 20. Nominal charges are measured in each endpoint year’s dollars; real growth uses December CPI-U. SD is the sample standard deviation.')
keys=['Unadjusted','Primary','Base only'];rows=[]
for c,label in [('intercept','Intercept'),('x','Account growth (log points)'),('ln_n19','Log baseline accounts')]:
 rr=[label]
 for k in keys:
  v=fits[k]['coefficients'].get(c);rr.append('—' if v is None else f"{v['estimate']:.3f}\n({v['HC3_SE']:.3f})")
 rows.append(rr)
rows += [['N']+[str(fits[k]['n']) for k in keys],['R-squared']+[f"{fits[k]['r2']:.3f}" for k in keys],['Account-growth 95% CI']+[f"[{fits[k]['ci_low']:.3f}, {fits[k]['ci_high']:.3f}]" for k in keys]]
tables['TABLE3']=('Table 3. Charge-growth regressions on the same 20 districts',['Variable','Unadjusted full charge','Primary full charge','Base only, same sample'],rows,'Outcome: 100 × log(B2024/B2019), at 4,000 gallons. Parentheses contain HC3 standard errors; intervals use t critical values. No significance stars. Full charge includes supported recurring extras.')
keys=['Broader supported','Omit Hyden-Leslie','Add log initial charge','Original design weights','Exclude mixed; ARC adjusted','Martin conditional scenario','2000 gallons','6000 gallons']
rows=[[k,str(fits[k]['n']),f"{fits[k]['slope']:.3f}",f"[{fits[k]['ci_low']:.3f}, {fits[k]['ci_high']:.3f}]"] for k in keys]
tables['TABLE4']=('Table 4. Selected sensitivity estimates',['Specification','N','Slope','95% HC3 interval'],rows,'Each row adjusts for baseline size. Except consumption sensitivities, outcomes use 4,000 gallons. Different samples do not identify the causal effect of exclusions.')
short=lambda n:n.replace(' County Water and Sewer District',' Co.').replace(' County Water & Sewer District',' Co.').replace(' County Water District',' Co.').replace(' Water District','').replace('Cumberland Falls Highway','Cumberland Falls Hwy.')
tables['TABLEA1']=('Table A1. Main district observations',['District (PSC ID)','Accounts 2019','Accounts 2024','Account change %','2019 charge','2024 charge','Charge growth %'],[[short(x.utility_name)+' ('+x.utility_id+')',f"{x['2019']:,.0f}",f"{x['2024']:,.0f}",f'{x.account_pct:.2f}',f"${x.charge_inclusive_bill_or_scenario_2019:.2f}",f"${x.charge_inclusive_bill_or_scenario_2024:.2f}",f"{x.full_pct_change:.2f}"] for _,x in m.sort_values('utility_name').iterrows()],'Source: reconstructed prices and PSC residential-account reports. Charges are nominal monthly dollars at 4,000 gallons; growth columns are ordinary percentages.')
rec=pd.read_csv(R/'results/Reconstruction_Summary_20_v4.csv',dtype={'utility_id':str})
rows=[[short(x.utility_name)+' ('+x.utility_id+')']+[f"{getattr(x,c):.2f}" for c in ['base_2019','extra_2019','base_2024','extra_2024','delta_base','delta_total']] for x in rec.itertuples()]
rows.append(['Sum of benchmark charges']+[f"{rec[c].sum():.2f}" for c in ['base_2019','extra_2019','base_2024','extra_2024','delta_base','delta_total']])
tables['TABLEB1']=('Table B1. Components of each standardized approved charge',['District (PSC ID)','2019 base','2019 extra','2024 base','2024 extra','Base change','Total change'],rows,'Nominal monthly dollars at 4,000 gallons; sums represent one hypothetical connection per district, not revenues. Net extras rise $7.33. Zero endpoint extras allow temporary charges. Sources: tariff register and event ledger.')
rows=[[short(x.utility_name)+' ('+x.utility_id+')',x.documented_proceedings,x.document_supported_interpretation,x.material_limit+' Sources: '+x.evidence_reference+'.'] for x in rec.itertuples()]
tables['TABLEB2']=('Table B2. Documentary interpretation and limits for all 20 districts',['District (PSC ID)','Documented regulatory route','Supported interpretation','Material limit and evidence reference'],rows,'PWA = purchased-water adjustment. Event and surcharge IDs link to resolved original sources in the replication ledger. Dates and monetary changes describe authorizations or benchmark calculations unless billing evidence is expressly stated. All districts lack a complete independently reconciled customer-invoice history.')
rendered_md=md
for key,(title,headers,rows,note) in tables.items():
 tab='\n'+title+'\n\n|'+'|'.join(headers)+'|\n|'+'|'.join(['---']*len(headers))+'|\n'+'\n'.join('|'+'|'.join(str(v).replace('\n','; ') for v in row)+'|' for row in rows)+'\n\n'+note+'\n'
 rendered_md=rendered_md.replace('{{'+key+'}}',tab)
(R/'paper/manuscript_complete.md').write_text(rendered_md.replace('{{EQUATION}}','yᵢ = α + βxᵢ + γ ln(Nᵢ,2019) + εᵢ').replace('{{FIGURE1}}','![Figure 1. Account change and standardized charge growth](../results/figures/figure1_association.png)').replace('{{FIGURE2}}','![Figure 2. Sensitivity estimates and uncertainty](../results/figures/figure2_sensitivity.png)'))

def para(txt,style=None):
 p=d.add_paragraph(style=style)
 # Native hyperlink for URLs; preserve scholarly author-date text.
 for part in re.split(r'(https?://\S+)',txt):
  if part.startswith('http'):
   rel=d.part.relate_to(part,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);run=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text=part;run.append(tx);h.append(run);p._p.append(h)
  else:p.add_run(part)
 return p

def table(key):
 title,headers,rows,note=tables[key];p=para(title);p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 widths={'TABLE1':[4.1,1.2,1.2],'TABLE2':[2.4,.82,.82,.82,.82,.82],'TABLE3':[2.0,1.5,1.5,1.5],'TABLE4':[3.25,.4,.7,2.15],'TABLEA1':[1.7,.8,.8,.8,.8,.8,.8],'TABLE5':[2.4,.85,.85,2.4],'TABLEB1':[1.9,.76,.76,.76,.76,.78,.78],'TABLEB2':[1.25,1.4,1.7,2.15]}[key]
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for row_idx,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trPr.append(ns)
  for i,c in enumerate(row.cells):
   c.width=Inches(widths[i])
   for p in c.paragraphs:
    p.paragraph_format.keep_with_next=(key in ['TABLE1','TABLE2','TABLE3','TABLE4'] or row_idx==0 or row_idx==len(t.rows)-1)
    p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(2 if key=='TABLEB1' else 3 if key in ['TABLEA1','TABLE2'] else 5);p.paragraph_format.space_before=Pt(2 if key=='TABLEB1' else 3 if key in ['TABLEA1','TABLE2'] else 5)
    for run in p.runs:run.font.size=Pt(10)
 for i,col in enumerate(t.columns):col.width=Inches(widths[i])
 t.autofit=False
 for rowidx,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   tcpr=cell._tc.get_or_add_tcPr()
   margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    node=OxmlElement('w:'+side);node.set(qn('w:w'),'40' if key in ['TABLEA1','TABLEB1','TABLE2'] else '65');node.set(qn('w:type'),'dxa');margins.append(node)
   tcpr.append(margins)
   valign=OxmlElement('w:vAlign');valign.set(qn('w:val'),'center');tcpr.append(valign)
   if rowidx==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');tcpr.append(sh)
    for pp in cell.paragraphs:
     for rr in pp.runs:rr.bold=True
   if key in ['TABLEB1','TABLE2','TABLE3','TABLEA1'] and ci>0:
    for pp in cell.paragraphs:pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  node=OxmlElement('w:'+side);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
 t._tbl.tblPr.append(borders)
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 p=para(note);p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_before=Pt(6)
 for run in p.runs:run.font.size=Pt(10)

in_references=False
for block in md.split('\n\n'):
 block=block.strip()
 if not block:continue
 key=block.removeprefix('{{').removesuffix('}}')
 if key in tables:table(key)
 elif key in ['FIGURE1','FIGURE2']:
  name='figure1_association' if key=='FIGURE1' else 'figure2_sensitivity';p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(R/'results/figures'/f'{name}.png'),width=Inches(6.45))
  p=para('Figure 1. Residential-account change and standardized charge growth. Each point represents one district. The line is an unadjusted fit.' if key=='FIGURE1' else 'Figure 2. Selected account-growth estimates and 95% HC3 intervals. Every displayed interval includes zero.');p.paragraph_format.line_spacing=1
  for run in p.runs:run.font.size=Pt(10)
 elif key=='EQUATION':
  p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;math=OxmlElement('m:oMath');
  def mrun(parent,txt):
   mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text=txt;mr.append(mt);parent.append(mr)
  mrun(math,'yᵢ = α + βxᵢ + γ ln(')
  sub=OxmlElement('m:sSub');base=OxmlElement('m:e');mrun(base,'N');sub.append(base);idx=OxmlElement('m:sub');mrun(idx,'i,2019');sub.append(idx);math.append(sub)
  mrun(math,') + εᵢ');p._p.append(math)
 elif block.startswith('# '):para(block[2:],'Title')
 elif block.startswith('## '):
  in_references=block[3:]=='References'
  p=para(block[3:],'Heading 1')
  if block[3:] in ['1. Introduction','References','Appendix A. District observations and interpretation','Appendix B. Full-sample charge reconstruction']:p.paragraph_format.page_break_before=True
 elif block.startswith('### '):para(block[4:],'Heading 2')
 else:
  p=para(block)
  if in_references:
   p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(5)
   for run in p.runs:run.font.size=Pt(11)
  elif 'https://' in block and not block.startswith('The '):p.paragraph_format.line_spacing=1.0
d.save(R/'paper/KEA_Manuscript_v4.docx')
(R/'results/manuscript_tables.json').write_text(json.dumps(tables,indent=2))
print('Words (markdown, tables excluded):',len(md.split()))
